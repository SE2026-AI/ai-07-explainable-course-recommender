"""Xuất mọi bảng và view ra 1 file Excel để XEM (sửa dữ liệu thì sửa CSV trong data/seed).

Chạy: python -m tools.export_excel   ->  data/xem_du_lieu.xlsx
"""
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from src.data.db.build import get_engine
from src.data.db.schema import metadata
from src.data.db.views import VIEWS


def main(url=None, out="data/xem_du_lieu.xlsx"):          # url=None -> DATABASE_URL
    engine = get_engine(url)
    names = [t.name for t in metadata.sorted_tables] + list(VIEWS)
    with engine.connect() as conn, pd.ExcelWriter(out, engine="openpyxl") as w:
        tong = []
        for n in names:
            df = pd.read_sql(text(f"SELECT * FROM {n}"), conn)
            tong.append({"Bảng/view": n, "Số dòng": len(df),
                         "Mô tả": (metadata.tables[n].comment if n in metadata.tables else "View") or ""})
            df.astype(str).replace({"None": "", "nan": "", "NaT": ""}).to_excel(w, sheet_name=n[:31], index=False)
        pd.DataFrame(tong).to_excel(w, sheet_name="_tong_quan", index=False)
        w.book.move_sheet("_tong_quan", offset=-len(names))
        for ws in w.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
            for i, col in enumerate(ws.columns, start=1):
                ws.column_dimensions[get_column_letter(i)].width = min(55, max(10, max(len(str(c.value or "")) for c in col) + 2))
    print(f"Đã tạo {out}")


if __name__ == "__main__":
    main()
